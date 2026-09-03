#!/usr/bin/env python3
"""Build bilingual academic-paper deliverables from frozen v2 evidence."""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DESKTOP = Path("/share/home/yuanwenjie/Desktop")
FIGDIR = DESKTOP / "图片"
OUT = DESKTOP / "OpenROAD_Evolve_学术论文_重制版"
ASSETS = OUT / "figures"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(exist_ok=True)

FIGS = [
    ("fig1.png", "01_完整workflow_重制版.png"),
    ("fig2.png", "02_Agent框架与执行约束_重制版.png"),
    ("fig3.png", "03_自演化学习机制_重制版.png"),
    ("fig4.png", "04_知识卡实例_重制版.png"),
    ("fig5.png", "05_BO_GP参数探索_重制版.png"),
    ("fig6.png", "06_EDA到AI数据接口_重制版.png"),
    ("fig7.png", "07_RTL生成链_重制版.png"),
    ("fig8.png", "08_论文实验结果_重制版.png"),
]
for dst, src in FIGS:
    shutil.copy2(FIGDIR / src, ASSETS / dst)
    shutil.copy2(FIGDIR / "English" / src.replace(".png", "_English.png"), ASSETS / dst.replace(".png", "_en.png"))


ZH = {
"title": "OpenROAD–Evolve：面向自然语言 RTL 与物理设计的受保护、证据驱动 Agentic EDA 闭环",
"abstract": "大语言模型能够调用 EDA 工具，但工具调用本身不能证明设计正确、优化有效或经验可迁移。本文提出 OpenROAD–Evolve：一个以 OpenROAD/OpenROAD-flow-scripts 为受保护执行后端的 Agentic EDA 平台。系统把自然语言前端、重复测量的多目标参数探索、低失真 EDA-to-AI 数据接口、受约束多 Agent 编排和可证伪知识准入统一在同一证据合同下。RTL 前端将规格解析、测试生成和 RTL 编写分离，并由 RTLScout 通过编译、lint、冻结 testbench 仿真、mutation adequacy 和成本门迭代候选。后端为面积、setup WNS 与功耗分别拟合 exact RBF Gaussian process，以标量化 expected improvement 与经验可行率提出联合参数向量，每个向量由三个配对 OR_SEED 复测。学习层采用 2×2 组合干预与未见设计 holdout，将经验标记为 validated 或 refuted，而不直接授予执行权。冻结实验包含 20 次独立 RTL 尝试、960 次参数探索 full-flow、288 次迁移复验 full-flow、240 个数据接口问答和 Agent 故障注入。RTL 完整通过率为 18/20；BO/GP 在 40 个配对单元中胜 33 个，平均配对效应为 0.001614（bootstrap 95% CI [0.000923, 0.002354]，sign-flip p=0.000490）；Typed EDAIR 问答准确率为 93.33%，而 KPI-only 为 4.17%；12 条迁移中 6 条 validated、6 条 refuted。结果证明该平台已形成可运行、可审计的 v2 全链路，但不证明四个 RTL 任务可代表任意芯片、BO/GP 对任意设计均显著，或数据接口与 Agent 安全性会自动提高 PPA。",
"sections": [
("1 引言", [
"Agentic EDA 的核心难题不是让模型输出一段脚本，而是把不确定的语言提案接到确定、昂贵且有强约束的芯片实现流程中。一次 RTL 生成可能语法正确却功能错误；一次参数修改可能只利用随机 seed 的波动；一次在 GCD 上有效的经验可能在 FIFO 上反向；一段被压缩的日志摘要可能丢掉精确对象、单位和来源位置。若模型同时写设计、写测试并宣布成功，系统就形成了不可审计的自证循环。",
"OpenROAD–Evolve 将上述问题统一为 protected evidence contract：Agent 只能产生结构化提案，Runtime 固定工具链、输入、预算、checkpoint 和 run ID，Reviewer 根据硬约束、重复测量与来源完整性决定晋级。本文贡献有四点：（1）建立自然语言到 SpecIR、双 Agent 验证和 RTLScout 候选演化的唯一 RTL 入口；（2）建立基于 exact RBF GP、EI 和真实 full-flow 复测的组合参数闭环；（3）建立保留原始 artifact 与 loss manifest 的 Typed EDAIR；（4）建立包含反证和 holdout 的知识状态机，并以五类消融分别验证功能、优化、数据、学习和编排。"
]),
("2 相关工作与研究定位", [
"AIVRIL2（DATE 2025）以协作 Agent 读取 EDA 错误并修复 RTL；VeriOpt（ICCAD 2025）使用 Planner、Programmer、Reviewer 和 Evaluator 进行 PPA-aware RTL 生成；AutoSilicon（TODAES 2025）强调层级任务分解；EvolVE（2026）研究 evolutionary search、MCTS、idea-guided refinement 与结构化 testbench。本文吸收角色分离与迭代验证，但将 testbench 冻结、mutation adequacy、候选 lineage 和后端独立复测放入统一合同。",
"ORFS-agent（MLCAD 2025）展示 tool-using agent 的芯片优化；AgenticPD（2026）强调 stage-aware physical-design agent；ReviewDSE（MLCAD 2026）把搜索边界从公共参数推进到受保护白盒机制。本文 v2 不执行 TimingECO、EvoDRC 或源码修改，而是先把参数级闭环、证据接口、停滞换向和知识准入打通。CircuitOps（ICCAD 2023）证明 cell/pin/net 关系表是 AI4EDA 基础设施；EDATracer（2026）进一步强调大规模 artifact 分析。Typed EDAIR 因而不以单段摘要替代原始材料，而提供 L0 KPI、L1 stage、L2 object graph 与 L3 raw excerpt。"
]),
("3 系统方法", [
"图1给出三平面架构。RTL 构造平面把自然语言转为 SpecIR，Verification Agent 独立生成 TB/SVA 与 mutation plan，RTL Author Agent 只修改候选 RTL。RTLScout 的 evaluate 循环顺序执行 compile/elaborate、lint/structural、冻结 TB 仿真、mutation adequacy 与 Yosys/ABC cost；RTL 错误反馈 Author，测试过弱反馈 Verifier。只有功能门全部通过的候选才进入独立 replay 和 OpenROAD baseline。",
"物理设计平面从固定 baseline 开始。对每个目标 m∈{area,WNS,power}，平台在归一化参数 x 上拟合 y_m(x)~GP(mu_m,k_RBF)，length-scale 固定为 0.35；重复配置先按均值聚合，噪声取 sample variance/n。偏好配置将三个目标变为 scalar utility，timing 与 performance 当前为同义 profile。下一候选为 512 点 deterministic Latin-hypercube 池中 EI(x)×p_feasible(x) 最大的未测向量。每个向量运行三个配对 OR_SEED；WNS≥0、DRC=0 等为硬门。连续三批没有达到预注册 0.5% 可靠改善时，系统冻结 trace，进入 stage diagnosis 和搜索换向，但 v2 不声称已执行 v3 repair 工具。",
"图2的八阶段链为 Map、Semantics、Experiment、Hypothesis、Implement、Validate、Review 和 Memory。中央 Runtime 是唯一执行权威。Observer 无执行权；Planner 只能输出 ExperimentPlan；Hypothesis 的 execution_allowed=false；ActionSpec 必须满足 schema、白名单与预算；Reviewer 决定晋级或拒绝。该设计使“思考过程”成为可检查的 typed products，而不是自由文本。",
"图3展示学习准入。对因素 A、B 运行四个组合并估计交互效应 Delta=y++-y+--y-++y--。来源设计上的方向只形成 candidate knowledge；同一干预必须在未见设计 holdout 上按同协议复验。方向一致且硬门通过进入 validated；方向翻转、效应不足或协议不兼容进入 refuted。两种知识都保存 context、falsifier、run IDs、artifact SHA 和 uncertainty，且 execution_allowed 始终为 false。",
"图6展示 Typed EDAIR。Artifact Registry 保存 log、report、RTL、netlist、SDC、Liberty、DEF、ODB、GDS 与 SPEF 的 SHA-256、tool commit、flow seed、parser version 和 source span。Parser 输出 metric、stage、instance、pin、net、timing path、physical coordinate、UNKNOWN 与 loss_manifest。Agent 可逐级查询，证据不足时以 artifact ID 回读原始字节，因此结构化不会切断原始来源。"
]),
("4 实验设计", [
"本文提出五个研究问题：RQ1，双 Agent+RTLScout 能否在固定跨类型任务上自动产生可晋级 RTL？RQ2，同预算 BO/GP 是否优于 seeded random？RQ3，Typed EDAIR 是否提高基于真实 artifact 的诊断问答正确率？RQ4，组合干预与 holdout 能否识别错误迁移？RQ5，checkpoint、authority gate 与 review gate 是否对可靠编排必要？",
"RTL 冻结矩阵包含 FIFO、GCD、ibex_alu 与 UART TX，每题 5 次独立尝试，共 20 次；ibex_alu 只是 ALU 子模块。参数实验共有 40 个 design×policy-seed 配对单元、960/960 次真实 ORFS runs。学习实验覆盖 12 个有向 source→holdout 对、288/288 次真实 ORFS runs。EDAIR 使用 240 个问题，比较 KPI-only 与 Typed EDAIR。Agent 消融删除 checkpoint、authority gate 或 review gate，并注入可判定故障。所有修复后工程回归与冻结主统计分开报告。"
]),
("5 结果", [
"RQ1：20 次 RTL 尝试中 18 次完整通过，首稿完整门通过率 65%，RTLScout 救回 5 次；FIFO 5/5、GCD 5/5、ibex_alu 4/5、UART TX 4/5。两个失败用例修复后独立回归 2/2，但不回写冻结的 18/20。该结果证明固定题库上的自动链路，不等于任意自然语言中小芯片能力。",
"RQ2：BO/GP 在 40 个配对单元中胜 33 个，seeded random 胜 7 个。平均配对差 0.0016138，中位差 0.0013741，bootstrap 95% CI [0.0009234,0.0023536]，paired sign-flip p=0.000490。达到 0.5% 阈值的比例为 52.5% 对 30%。然而逐设计 Holm 仅 GCD 与 ibex_alu 显著；四设计聚类敏感性 exact p=0.125，因此不能声称跨任意设计普遍显著。",
"RQ3：KPI-only 为 10/240（4.17%），UNKNOWN 92.08%，false answer 3.75%；Typed EDAIR 为 224/240（93.33%），UNKNOWN 5%，false answer 1.67%。调用级 exact p=3.81e-6，但四设计聚类敏感性 p=0.125。该实验验证信息可用性，不证明 PPA 因数据接口自动提高。",
"RQ4：12 条迁移中 6 条 validated、6 条 refuted，系统阻止 6 条错误迁移。真实卡片中，GCD 来源交互为 +1.862；GCD→UART TX 为 +0.532，进入 validated；GCD→FIFO 为 -1.596，方向翻转后进入 refuted。RQ5：完整架构出现 0 duplicate、0 unsupported executable hypothesis、0 below-threshold promotion，证据完整率 100%；删除 checkpoint 产生 2 次重复 run，删除 authority gate 使 12 个 hypothesis 失去不可执行保护，删除 review gate 会误晋级 8 个低于阈值的正波动。"
]),
("6 有效性威胁与结论", [
"外部有效性受四个 RTL 任务、有限 PDK/设计与共享工具链限制；OR_SEED 是重复测量，不是独立设计。内部有效性受固定 GP length-scale、标量化效用、自动 testbench 覆盖范围和 parser 正确性限制。构念有效性方面，Agent 安全违规为零不等于 QoR 提升，QA 正确率不等于 PPA 改善，mutation adequacy 不等价于形式证明。统计结论必须同时报告配对单元和设计聚类敏感性。",
"OpenROAD–Evolve v2 已把自然语言 RTL、真实 baseline、重复 BO/GP、停滞换向、Typed EDAIR、受约束 Agent 和反证学习接成一个可复现闭环。下一阶段可以把 TimingECO、Resynth、EvoDRC 与源码级工具作为受限 ActionSpec 插件接入，但这些动作必须继续经过相同的 protected Runtime、fresh extraction、硬门和 holdout 证据合同。"
])],
}

EN = {
"title": "OpenROAD–Evolve: A Protected, Evidence-Grounded Agentic EDA Loop from Natural-Language RTL to Physical Design",
"abstract": "Large language models can invoke EDA tools, but invocation alone does not establish functional correctness, optimization efficacy, or transferable learning. We present OpenROAD–Evolve, an agentic EDA platform that uses OpenROAD/OpenROAD-flow-scripts as a protected execution backend. It unifies a natural-language RTL front end, replicated multi-objective parameter exploration, a low-loss EDA-to-AI interface, constrained multi-agent orchestration, and falsifiable knowledge admission under one evidence contract. Specification parsing, verification generation, and RTL authoring are separated; RTLScout evolves candidates through compilation, lint, frozen-testbench simulation, mutation adequacy, and cost gates. For physical design, independent exact RBF Gaussian processes model area, setup WNS, and power. Scalarized expected improvement multiplied by empirical feasibility selects joint parameter vectors, each replayed with three paired OR_SEED values. Learning uses 2x2 interventions and unseen-design holdouts, storing both validated and refuted knowledge without granting execution authority. Frozen evaluation covers 20 independent RTL attempts, 960 parameter-search full flows, 288 transfer-validation full flows, 240 data-interface questions, and injected agent failures. RTL passes 18/20 attempts. BO/GP wins 33 of 40 paired cells with mean paired effect 0.001614 (bootstrap 95% CI [0.000923, 0.002354], sign-flip p=0.000490). Typed EDAIR reaches 93.33% diagnostic QA accuracy versus 4.17% for KPI-only context. Six of 12 transfers are validated and six refuted. These results support an executable and auditable v2 loop, but do not imply arbitrary-chip RTL generalization, universal BO superiority, or automatic PPA gains from data and safety infrastructure.",
"sections": [
("1 Introduction", [
"Agentic EDA must connect probabilistic language proposals to deterministic, expensive, and strongly constrained implementation flows. Syntactically valid RTL can be functionally wrong; a favorable parameter result can be seed noise; a GCD interaction can reverse on FIFO; and a compressed log summary can omit the precise object, unit, or source location required for diagnosis. Allowing the same model to write a design, write its test, and declare success creates an unauditable self-confirmation loop.",
"OpenROAD–Evolve treats these failures as one systems problem: a protected evidence contract. Agents emit typed proposals. Runtime fixes tools, inputs, budgets, checkpoints, run identifiers, and artifacts. Reviewer gates promotion using hard constraints, replicated measurements, and provenance. Our contributions are: (1) one RTL entry path from natural language through SpecIR, independent verification, and RTLScout evolution; (2) joint-parameter BO/GP driven by real replicated full flows; (3) Typed EDAIR with raw-artifact fallback and an explicit loss manifest; and (4) a knowledge state machine that preserves refutation and tests transfer on holdouts."
]),
("2 Related Work and Positioning", [
"AIVRIL2 (DATE 2025) uses collaborative agents and EDA diagnostics for RTL repair; VeriOpt (ICCAD 2025) studies planner/programmer/reviewer/evaluator roles for PPA-aware Verilog; AutoSilicon (TODAES 2025) scales RTL generation through hierarchical decomposition; EvolVE (2026) studies evolutionary search, MCTS, idea-guided refinement, and structured testbench generation. We adopt role separation and iterative feedback while adding frozen verification packages, mutation adequacy, candidate lineage, and independent backend replay.",
"ORFS-agent (MLCAD 2025) demonstrates tool-using chip-design optimization, AgenticPD (2026) emphasizes stage-aware physical-design agents, and ReviewDSE (MLCAD 2026) explores protected white-box mechanisms. Our v2 scope intentionally stops before TimingECO, EvoDRC, or source edits: it first establishes a parameter-level closed loop, a protected evaluator, stall redirection, and evidence admission. CircuitOps (ICCAD 2023) motivates relational cell/pin/net infrastructure; EDATracer (2026) emphasizes artifact-scale analysis. Typed EDAIR therefore preserves source bytes instead of replacing them with one summary."
]),
("3 Method", [
"The RTL plane parses a natural-language request into SpecIR. A Verification Agent generates TB/SVA and a mutation plan, while an RTL Author can edit only candidate RTL. RTLScout evaluates compile/elaboration, lint/structure, frozen-testbench simulation, mutation adequacy, and Yosys/ABC cost. RTL failures return to the author; weak tests return to the verifier. Only candidates passing every functional gate enter independent replay and OpenROAD baseline measurement.",
"The physical-design plane begins with a fixed baseline. For each objective m in {area, WNS, power}, it fits y_m(x)~GP(mu_m,k_RBF) on normalized joint parameters with fixed length scale 0.35. Repeated points are aggregated by their mean; observation noise is sample variance divided by replica count. A user preference scalarizes the objectives; timing and performance are currently aliases. The next untested vector maximizes EI(x) times local empirical feasibility over a deterministic 512-point Latin-hypercube pool. Every vector is evaluated with three paired OR_SEED replicas. WNS>=0 and DRC=0 remain hard constraints. Three consecutive batches without a preregistered 0.5% reliable gain freeze the trace and trigger stage diagnosis and search redirection; v2 does not claim execution of v3 repair tools.",
"Eight typed stages—Map, Semantics, Experiment, Hypothesis, Implement, Validate, Review, and Memory—surround the only execution authority, Runtime. Observer cannot execute; Planner emits only ExperimentPlan; hypotheses carry execution_allowed=false; ActionSpec requires schema, whitelist, and budget; Reviewer alone promotes or rejects.",
"Learning uses four A/B combinations and interaction Delta=y++-y+--y-++y--. A source effect creates candidate knowledge only. The same intervention is replayed on an unseen holdout. Directional agreement plus hard-gate success yields validated; sign reversal, insufficient effect, or incompatibility yields refuted. Both states preserve context, falsifier, uncertainty, run IDs, and artifact hashes, and neither grants execution authority.",
"Typed EDAIR registers logs, reports, RTL, netlists, SDC, Liberty, DEF, ODB, GDS, and SPEF with SHA-256, tool commit, flow seed, parser version, and source span. Parsers emit metrics, stages, instances, pins, nets, timing paths, coordinates, UNKNOWN, and loss_manifest. Agents query L0 KPI, L1 stage, L2 object graph, or L3 raw excerpt and can read exact source bytes by artifact ID."
]),
("4 Experimental Design", [
"We ask five questions: RQ1 whether separated agents plus RTLScout produce promotable RTL across a fixed heterogeneous suite; RQ2 whether equal-budget BO/GP outperforms seeded random; RQ3 whether Typed EDAIR improves source-grounded diagnostic QA; RQ4 whether intervention plus holdout detects negative transfer; and RQ5 whether checkpoint, authority, and review gates are necessary.",
"The frozen RTL matrix has FIFO, GCD, ibex_alu, and UART TX, five independent attempts each; ibex_alu is an ALU submodule. Parameter evaluation contains 40 design-by-policy-seed paired cells and 960/960 real ORFS runs. Learning contains 12 directed source-to-holdout pairs and 288/288 real ORFS runs. EDAIR compares KPI-only and typed context on 240 questions. Agent ablations remove checkpoint, authority, or review gates. Post-freeze engineering regressions are reported separately from primary statistics."
]),
("5 Results", [
"RQ1: 18 of 20 RTL attempts pass the complete flow. First-candidate complete-gate pass rate is 65%, and RTLScout rescues five attempts: FIFO 5/5, GCD 5/5, ibex_alu 4/5, and UART TX 4/5. Two retained failures pass a later 2/2 engineering regression, which does not rewrite the frozen 18/20 result.",
"RQ2: BO/GP wins 33 of 40 paired cells and seeded random wins seven. Mean paired difference is 0.0016138, median 0.0013741, bootstrap 95% CI [0.0009234, 0.0023536], paired sign-flip p=0.000490. The 0.5% threshold is reached in 52.5% versus 30% of cells. Holm-adjusted per-design tests are significant only for GCD and ibex_alu; four-design clustered sensitivity gives exact p=0.125.",
"RQ3: KPI-only answers 10/240 correctly (4.17%), with 92.08% UNKNOWN and 3.75% false answers. Typed EDAIR answers 224/240 correctly (93.33%), with 5% UNKNOWN and 1.67% false answers. Invocation-level exact p=3.81e-6, while four-design clustered sensitivity is p=0.125. This establishes information usability, not downstream PPA improvement.",
"RQ4: six of 12 transfers are validated and six refuted, blocking six erroneous transfers. The GCD source interaction is +1.862; GCD-to-UART-TX is +0.532 and validated; GCD-to-FIFO is -1.596 and refuted. RQ5: the full architecture has zero duplicates, unsupported executable hypotheses, and sub-threshold promotions, with 100% evidence completeness. Removing checkpoint creates two duplicate runs; removing authority exposes 12 unsafe hypotheses; removing review falsely promotes eight sub-threshold fluctuations."
]),
("6 Threats to Validity and Conclusion", [
"External validity is limited by four RTL tasks, a small design/PDK set, and a shared toolchain; OR_SEED is a replica, not an independent design. Internal validity is affected by the fixed GP length scale, scalar utility, generated-test coverage, and parser correctness. Construct validity requires care: zero agent-safety violations is not QoR gain, QA accuracy is not PPA gain, and mutation adequacy is not formal proof. Statistical claims must retain both paired-cell and design-cluster views.",
"OpenROAD–Evolve v2 connects natural-language RTL, real baselines, replicated BO/GP, stall redirection, Typed EDAIR, constrained agents, and refutation-aware learning into one reproducible loop. Future TimingECO, Resynth, EvoDRC, and source-level tools should enter only as bounded ActionSpec plugins under the same protected Runtime, fresh extraction, hard gates, and holdout evidence contract."
])],
}

REFS = [
"Ghose et al. ORFS-agent: Tool-Using Agents for Chip Design Optimization. MLCAD 2025. DOI:10.1109/MLCAD65511.2025.11189204.",
"Islam et al. EDA-Aware RTL Generation with Large Language Models (AIVRIL2). DATE 2025. DOI:10.23919/DATE64628.2025.10992789.",
"Tasnia et al. VeriOpt: PPA-Aware High-Quality Verilog Generation via Multi-Role LLMs. ICCAD 2025. DOI:10.1109/ICCAD66269.2025.11240771.",
"Li et al. AutoSilicon: Scaling Up RTL Design Generation Capability of Large Language Models. ACM TODAES 30(6), 2025. DOI:10.1145/3737286.",
"Hsin et al. EvolVE: Evolutionary Search for LLM-based Verilog Generation and Optimization. arXiv:2601.18067, 2026.",
"Ren et al. AgenticPD: A Stage-Aware Agentic Framework for Physical Design QoR Optimization. arXiv:2607.04758v2, 2026.",
"Tieu et al. EDATracer: An Agentic Framework for Large-Scale EDA Artifact Analysis. arXiv:2608.04032, 2026.",
"Zheng et al. From Tool Invocation to Source-Mechanism Exploration: Protected White-Box DSE for Open-Source EDA. MLCAD 2026, arXiv:2607.11294v4.",
"Liang et al. CircuitOps: An ML Infrastructure Enabling Generative AI for VLSI Circuit Optimization. ICCAD 2023. DOI:10.1109/ICCAD57390.2023.10323611.",
"Wu et al. EvoDRC: A Self-Evolving Agentic Framework for Automated DRC Violation Repair. arXiv:2607.20019, 2026.",
]


def tex_escape(s: str) -> str:
    for a,b in [("\\","\\textbackslash{}"),("&","\\&"),("%","\\%"),("#","\\#"),("_","\\_"),("≥","$\\ge$"),("→","$\\rightarrow$")]: s=s.replace(a,b)
    return s


def build_tex(data, language):
    chinese = language == "zh"
    pre = r'''\documentclass[10pt,twocolumn]{article}
\usepackage[a4paper,margin=16mm,columnsep=7mm]{geometry}
\usepackage{fontspec}\usepackage{graphicx}\usepackage{xcolor}\usepackage{booktabs}\usepackage{amsmath}\usepackage{hyperref}\usepackage{caption}
\setmainfont{''' + ("SimHei" if chinese else "DejaVu Serif") + r'''}
\setsansfont{''' + ("SimHei" if chinese else "DejaVu Sans") + r'''}
\definecolor{navy}{HTML}{174A7E}\definecolor{orange}{HTML}{E47A36}
\hypersetup{colorlinks=true,linkcolor=navy,urlcolor=navy}\captionsetup{font=small,labelfont=bf}
''' + (r'''\renewcommand{\abstractname}{摘要}\renewcommand{\figurename}{图}
''' if chinese else "") + r'''
\title{\textbf{''' + tex_escape(data["title"]) + r'''}}
\author{Yuan Wenjie et al.\\\small Author order, affiliation, corresponding author, and funding to be confirmed}
\date{August 2026}
\begin{document}\maketitle
\begin{abstract}''' + tex_escape(data["abstract"]) + r'''\end{abstract}
\noindent\textbf{Keywords:} Agentic EDA; OpenROAD; RTL generation; Bayesian optimization; EDA-to-AI; self-evolution.
'''
    body=[]
    figidx=0
    captions_zh=["完整系统工作流与受保护执行边界。","八阶段 Agent 链与 Runtime 权限核。","组合干预、holdout 与知识状态机。","validated 与 refuted 知识卡实例。","联合参数空间中的 GP 后验与 BO 迭代。","保留 provenance 和原始回读的 Typed EDAIR。","双 Agent 隔离与 RTLScout 候选演化。","五个研究问题的冻结结果总览。"]
    captions_en=["Complete workflow and protected execution boundary.","Eight-stage agent chain around Runtime authority.","Intervention, holdout, and knowledge-state admission.","Validated and refuted evidence-card examples.","GP posterior and BO iteration over a joint parameter space.","Typed EDAIR with provenance and raw-artifact fallback.","Separated agents and RTLScout candidate evolution.","Frozen results for five research questions."]
    inserts={"3 系统方法":[0,1,2,5,6],"5 结果":[3,4,7],"3 Method":[0,1,2,5,6],"5 Results":[3,4,7]}
    for title,paras in data["sections"]:
        body.append("\\section{"+tex_escape(title)+"}")
        for p in paras: body.append(tex_escape(p)+"\n\n")
        for idx in inserts.get(title,[]):
            cap=(captions_zh if chinese else captions_en)[idx]
            suffix="" if chinese else "_en"
            body.append(f'''\\begin{{figure*}}[t]\n\\centering\\includegraphics[width=0.98\\textwidth]{{figures/fig{idx+1}{suffix}.png}}\n\\caption{{{tex_escape(cap)}}}\\label{{fig:f{idx+1}}}\n\\end{{figure*}}''')
    body.append("\\section*{"+("参考文献" if chinese else "References")+"}")
    body.append("\\begin{enumerate}")
    body.extend("\\item "+tex_escape(r) for r in REFS)
    body.append("\\end{enumerate}")
    body.append("\\section*{"+("可复现性声明" if chinese else "Reproducibility Statement")+"}")
    body.append(tex_escape("All reported primary numbers are generated from frozen analysis JSON files under artifacts/v2-paper-*-20260825. Post-freeze engineering regressions are not merged into primary statistics. The repository records protocol, tool version, run identifiers, artifact hashes, and analysis scripts."))
    return pre+"\n".join(body)+"\n\\end{document}\n"


def add_field(p, text, bold=False):
    r=p.add_run(text);r.bold=bold;r.font.size=Pt(10.5);r.font.name="SimHei";r._element.rPr.rFonts.set(qn("w:eastAsia"),"SimHei")


def build_docx(data, lang, out):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    styles=d.styles
    for name in ["Normal","Title","Heading 1","Heading 2"]:
        styles[name].font.name="SimHei" if lang=="zh" else "Arial";styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"),"SimHei")
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_field(p,data["title"],True);p.runs[0].font.size=Pt(18)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_field(p,"Yuan Wenjie et al. · author metadata to be confirmed")
    d.add_heading("摘要" if lang=="zh" else "Abstract",level=1);add_field(d.add_paragraph(),data["abstract"])
    inserts={"3 系统方法":[0,1,2,5,6],"5 结果":[3,4,7],"3 Method":[0,1,2,5,6],"5 Results":[3,4,7]}
    for title,paras in data["sections"]:
        d.add_heading(title,level=1)
        for text in paras:
            p=d.add_paragraph();p.paragraph_format.first_line_indent=Inches(.25);p.paragraph_format.line_spacing=1.18;add_field(p,text)
        for idx in inserts.get(title,[]):
            suffix="" if lang=="zh" else "_en"
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ASSETS/f"fig{idx+1}{suffix}.png"),width=Inches(6.9))
            c=d.add_paragraph(f"Figure {idx+1}. "+(["完整工作流","Agent 框架","自演化学习","知识卡实例","BO/GP 参数探索","EDA-to-AI 接口","RTL 生成链","实验结果总览"][idx] if lang=="zh" else ["Complete workflow","Agent architecture","Self-evolution","Knowledge cards","BO/GP exploration","EDA-to-AI interface","RTL generation","Experimental results"][idx]));c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("参考文献" if lang=="zh" else "References",level=1)
    for i,r in enumerate(REFS,1):add_field(d.add_paragraph(),f"[{i}] {r}")
    d.add_heading("可复现性声明" if lang=="zh" else "Reproducibility Statement",level=1);add_field(d.add_paragraph(),"Primary results come from frozen artifacts/v2-paper-*-20260825 analyses; post-freeze regressions are reported separately.")
    d.save(out)


def main():
    files=[]
    for lang,data,stem in [("zh",ZH,"OpenROAD_Evolve_学术论文_中文版_重制版"),("en",EN,"OpenROAD_Evolve_Academic_Paper_English_Rebuilt")]:
        tex=OUT/f"{stem}.tex";tex.write_text(build_tex(data,lang),encoding="utf-8");files.append(tex)
        build_docx(data,lang,OUT/f"{stem}.docx")
        for _ in range(2): subprocess.run(["lualatex","-interaction=nonstopmode","-halt-on-error",tex.name],cwd=OUT,check=True,stdout=subprocess.DEVNULL)
    for ext in ("aux","log","out"):
        for p in OUT.glob(f"*.{ext}"):p.unlink()
    (OUT/"README.md").write_text("重制版包含中英文 PDF、DOCX、LaTeX 源文件及论文内嵌 figures。作者、单位、通讯作者和基金信息为投稿前人工确认项。\n",encoding="utf-8")


if __name__=="__main__": main()
